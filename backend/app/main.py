import io
import re
import os
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from lxml.etree import XMLSyntaxError
from openai import OpenAIError
from pydantic import ValidationError
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded
from slowapi.util import get_remote_address

from app.analyzer import AnalysisProviderError, analyze, improve_section
from app.schemas import (
    AnalysisRequest,
    AnalysisResponse,
    ImprovementRequest,
    ImprovementResponse,
    ResumeExportRequest,
    RoleRecommendation,
    SkillGap,
    LearningSuggestion,
)

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

limiter = Limiter(key_func=get_remote_address)
app = FastAPI(title="CareerLens AI API", version="1.0.0")
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
allowed_origins = [
    origin.strip()
    for origin in os.getenv(
        "FRONTEND_ORIGIN",
        "http://localhost:5173,http://127.0.0.1:5173,http://localhost:5174,http://127.0.0.1:5174",
    ).split(",")
    if origin.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type"],
)


def extract_resume(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Upload a PDF or DOCX file.")
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File size must be 5 MB or less.")
    try:
        if extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise HTTPException(status_code=400, detail="The file does not appear to be a valid PDF.")
            with fitz.open(stream=content, filetype="pdf") as document:
                if document.page_count > 100:
                    raise HTTPException(status_code=400, detail="The PDF must contain 100 pages or fewer.")
                text = "\n".join(page.get_text() for page in document)
        else:
            if not zipfile.is_zipfile(io.BytesIO(content)):
                raise HTTPException(status_code=400, detail="The file does not appear to be a valid DOCX document.")
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise HTTPException(status_code=400, detail="The file does not contain a readable DOCX document.")
                if sum(member.file_size for member in archive.infolist()) > 20 * 1024 * 1024:
                    raise HTTPException(status_code=413, detail="The extracted DOCX content is too large to process.")
            document = Document(io.BytesIO(content))
            text = "\n".join([paragraph.text for paragraph in document.paragraphs]
                             + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    except HTTPException:
        raise
    except (
        fitz.FileDataError,
        zipfile.BadZipFile,
        PackageNotFoundError,
        XMLSyntaxError,
        ValueError,
        OSError,
        RuntimeError,
        KeyError,
    ):
        raise HTTPException(status_code=400, detail="The file could not be read. Try another PDF or DOCX.")
    text = "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines() if line.strip())
    if len(text) < 30:
        raise HTTPException(status_code=400, detail="No readable resume text was found. Try a text-based PDF or paste the text.")
    if len(text) > 30000:
        raise HTTPException(status_code=413, detail="Extracted resume text exceeds the 30,000 character limit.")
    return text


@app.get("/api/health")
def health():
    return {"status": "ok", "ai_enabled": os.getenv("AI_ENABLED", "false").lower() == "true"}


def _analyze_request(payload: AnalysisRequest) -> AnalysisResponse:
    try:
        return analyze(payload.resume_text, payload.target_role)
    except ValueError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except (OpenAIError, AnalysisProviderError, ValidationError, KeyError, TypeError) as exc:
        raise HTTPException(
            status_code=502,
            detail="AI analysis failed or returned invalid data. Try again or switch to demo mode.",
        ) from exc


@app.post("/api/resume/upload")
@limiter.limit("20/minute")
async def upload_resume(request: Request, file: UploadFile = File(...)):
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    text = extract_resume(file.filename or "", content)
    return {"filename": Path(file.filename or "resume").name, "text": text}


@app.post("/api/resume/analyze", response_model=AnalysisResponse)
@limiter.limit("10/minute")
def analyze_resume(request: Request, payload: AnalysisRequest):
    return _analyze_request(payload)


@app.post("/api/jobs/recommend", response_model=list[RoleRecommendation])
@limiter.limit("10/minute")
def recommend_jobs(request: Request, payload: AnalysisRequest):
    return _analyze_request(payload).job_recommendations


@app.post("/api/skills/gaps", response_model=list[SkillGap])
@limiter.limit("10/minute")
def skill_gaps(request: Request, payload: AnalysisRequest):
    return _analyze_request(payload).missing_skills


@app.post("/api/learning/roadmap", response_model=list[LearningSuggestion])
@limiter.limit("10/minute")
def learning_roadmap(request: Request, payload: AnalysisRequest):
    return _analyze_request(payload).learning_suggestions


@app.post("/api/resume/improve", response_model=ImprovementResponse)
@limiter.limit("10/minute")
def improve_resume(request: Request, payload: ImprovementRequest):
    try:
        suggested = improve_section(payload.resume_text, payload.section, payload.current_text)
    except ValueError as exc:
        raise HTTPException(status_code=503, detail=str(exc))
    except (OpenAIError, AnalysisProviderError) as exc:
        raise HTTPException(status_code=502, detail="The improvement service failed. Please try again.") from exc
    return ImprovementResponse(
        section=payload.section, original=payload.current_text, suggested=suggested,
        note="Review and edit this suggestion before using it. It may need your verification.",
    )


@app.post("/api/resume/export")
def export_resume(payload: ResumeExportRequest):
    document = Document()
    document.add_heading("Resume draft", level=0)
    for line in payload.resume_text.splitlines():
        document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="careerlens-resume-draft.docx"'},
    )
